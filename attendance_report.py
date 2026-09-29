"""Reads the Institute's own master "Class Attendance" workbook -- one sheet
per medium (English/Hindi in this course), a date + session-label header per
column, Y/N/blank per student per session, and "N/A" columns marking
holidays/exam days that carry no session at all. This is the canonical
term-long record a professor keeps by hand-ish; it's a different file from
any per-cohort `out/<cohort>/attendance.xlsx` left behind by the retired scan
pipeline, and the two are not reconciled against each other here.

Layout (fixed by the Institute's template, not something we control):
  row 1  title + occasional column labels (holidays, exam windows, ...)
  row 2  the date for every real class-session column
  row 3  header: SR / Roll Number / [Email] / Name / % / then "C1", "C2", ...
         per session column, or "N/A" for a spacer column with no class --
         the Email column comes and goes between edits of this workbook
         (seen both with and without it across this term), so column
         positions are found by header text below rather than hardcoded;
         a fixed NAME_COL/FIRST_SESSION_COL silently reads the wrong cells
         the next time a column gets inserted or removed by hand.
  row 4+ one row per student

The "%" column in the workbook is itself a formula --
`Y / (Y + N) * 100` over cells actually filled in, ignoring blanks (a
session not yet held) and N/A columns (no session that day) -- recomputed
here in Python rather than trusted from the cached cell value, so a summary
built right after someone edits the sheet by hand doesn't need Excel to have
recalculated and resaved it first.
"""
from __future__ import annotations
from datetime import datetime

import openpyxl
import pandas as pd

DATE_ROW = 2
LABEL_ROW = 3
FIRST_DATA_ROW = 4

# The workbook's third mark. "E" is an *excused* absence - the student was
# away but produced a medical receipt or another accepted reason. It is not
# the same as "N" and must never be counted as one: a warning mail triggered
# by an excused absence is a mail to someone who did everything right.
#
# EXCUSED_MODE decides what it does to the percentage:
#   "exclude" (default) - the session drops out of that student's denominator
#                         entirely, so it neither helps nor hurts them.
#   "present"           - counted as attendance.
# Either way `Absent` counts only unexcused absences, which is the number the
# shortfall mail and the defaulter list are based on.
EXCUSED_MODE = "exclude"

# A student cell may also carry "N/A" - the same marker the workbook already
# uses in row 3 for a spacer column, here meaning "this class does not apply
# to this student": most often a late admission, whose first weeks happened
# before they were on the roll. Like a blank it counts towards nothing, but
# unlike a blank it is *deliberate*, and the health check needs to tell the
# two apart or every late joiner looks like a data-entry gap.
NOT_APPLICABLE = {"N/A", "NA", "-", "\u2013", "\u2014"}   # hyphen, en dash, em dash


def sheet_names(path: str) -> list[str]:
    return openpyxl.load_workbook(path, read_only=True).sheetnames


# The workbook gained a third sheet: both batches on one list, with a `Batch`
# column saying which medium each student belongs to. It is a *view* of the two
# medium sheets, not a third cohort - every student appears on it and on their
# own sheet - so anything that adds sheets together has to leave it out or
# count everyone twice.
def is_combined(path: str, sheet: str) -> bool:
    """True for a whole-course sheet: one that names a Batch per student."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return False
    ws = wb[sheet]
    # a read-only sheet reports no dimensions until it is walked, so read the
    # header row itself rather than trusting ws.max_column
    for row in ws.iter_rows(min_row=LABEL_ROW, max_row=LABEL_ROW, values_only=True):
        return any(v is not None and str(v).strip().lower() == "batch" for v in row)
    return False


def is_register(path: str, sheet: str) -> bool:
    """True for a sheet that actually holds marks: one with a `Roll Number`
    header on the label row. The workbook has grown non-register sheets before
    (a `Dashboard` of live figures, for one), and offering those as a batch to
    read is how a page ends up reporting on a summary block."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return False
    ws = wb[sheet]
    for row in ws.iter_rows(min_row=LABEL_ROW, max_row=LABEL_ROW, values_only=True):
        return any(v is not None and str(v).strip().lower() == "roll number" for v in row)
    return False


def register_sheets(path: str) -> list[str]:
    """Every sheet that holds marks, in workbook order - batch sheets and the
    combined one, never the Dashboard or any other non-register tab. This is
    what a sheet picker should offer."""
    return [s for s in sheet_names(path) if is_register(path, s)]


def cohort_sheets(path: str) -> list[str]:
    """The per-medium registers only, in workbook order: what you sum over."""
    return [s for s in sheet_names(path)
            if is_register(path, s) and not is_combined(path, s)]


def combined_sheets(path: str) -> list[str]:
    return [s for s in sheet_names(path) if is_register(path, s) and is_combined(path, s)]


def sheet_for_cohort(path: str, cohort: str) -> str | None:
    """The batch sheet a cohort folder reads from: `out/english` -> `English`.

    Matched by name, ignoring case, among the cohort sheets - the same rule the
    Overview page uses. The health check used to keep its own
    `{"english": "English", "hindi": "Hindi"}` table, so a course whose batches
    were called anything else failed "Register sheet present" out of the box.
    """
    want = str(cohort).strip().lower()
    return next((s for s in cohort_sheets(path) if s.strip().lower() == want), None)


# The workbook also carries a hand-built `Dashboard` sheet: the course's own
# live view of the same marks (threshold, per-batch summary, attendance bands,
# and the list below the threshold). It is not a register - nothing is read
# *from* it for a mail or a seat - but it is what the course looks at, so the
# toolkit reads its settings rather than hardcoding a second copy of them, and
# the health check compares its figures against the register they come from.
DASHBOARD_SHEET = "Dashboard"
BANDS = ["< 50%", "50-75%", "75-90%", "90-100%"]


def _sheet_grid(path: str, sheet: str, max_col: int = 8) -> list[list]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    return [list(row) for row in wb[sheet].iter_rows(max_col=max_col, values_only=True)]


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def dashboard_blocks(path: str, sheet: str = DASHBOARD_SHEET) -> dict:
    """The course's own Dashboard sheet, read back as
    `{"settings": {...}, "summary": {metric: {column: value}}, "bands": {...}}`.

    Labels are taken as the sheet writes them, lower-cased. A sheet that isn't
    there, or has been rearranged, gives `{}` or the blocks it can still find -
    this is a convenience, and nothing downstream may depend on it existing.
    """
    grid = _sheet_grid(path, sheet)
    if not grid:
        return {}
    out: dict = {"settings": {}, "summary": {}, "bands": {}, "columns": []}
    mode = None
    for row in grid:
        head = str(row[0]).strip() if row and row[0] is not None else ""
        key = head.lower()
        if key == "settings":
            mode = "settings"
            continue
        if key == "summary":
            mode = "summary"
            out["columns"] = [str(v).strip() for v in row[1:] if v is not None]
            continue
        if key.startswith("attendance band"):
            mode = "bands"
            continue
        if key.startswith("students below"):
            mode = None
            continue
        if not head:
            continue
        if mode == "settings":
            out["settings"][key] = row[1]
        elif mode == "summary":
            vals = [row[i + 1] for i in range(len(out["columns"]))]
            out["summary"][key] = dict(zip(out["columns"], vals))
        elif mode == "bands":
            out["bands"][head] = _num(row[1])
    return out


def _header_cols(ws) -> dict[str, int]:
    """Header text (lowercased, stripped) -> column, from LABEL_ROW. Session
    columns ("C1", "C2", ...) aren't included here -- see `_session_columns`."""
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=LABEL_ROW, column=c).value
        if v is not None:
            out[str(v).strip().lower()] = c
    return out


def _session_columns(ws, first_session_col: int) -> list[tuple[int, datetime, str]]:
    """(column, date, label) for every column that's an actual class
    session -- has both a date (row 2) and a non-"N/A" label (row 3)."""
    cols = []
    for c in range(first_session_col, ws.max_column + 1):
        d = ws.cell(row=DATE_ROW, column=c).value
        label = ws.cell(row=LABEL_ROW, column=c).value
        if isinstance(d, datetime) and label and str(label).strip().upper() != "N/A":
            cols.append((c, d, str(label).strip()))
    return cols


def load(path: str, sheet: str, excused: str = EXCUSED_MODE) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Returns (students, sessions).

    students: Roll, Name, Batch, Present, Absent, Excused, NotApplicable,
    Held, Percent -- one row per student. `Batch` comes from the sheet's own
    `Batch` column where it has one (the combined sheet) and is blank on a
    per-medium sheet, where the sheet name already says it. `Absent` is unexcused absences only ("N"); `Excused` counts
    "E" marks; `Held` is the number of sessions that count towards this
    student's percentage, which under the default `excused="exclude"` leaves
    out their excused ones. Percent is None for a student with no session held
    yet (Held == 0), never a divide-by-zero 0%.

    sessions: Session (the "Cn" label), Date, Present, Absent, Percent --
    one row per class actually held so far, in date order. Useful for a
    term-progress chart.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    hdr = _header_cols(ws)
    try:
        roll_col = hdr["roll number"]
        name_col = hdr["name"]
        pct_col = hdr["%"] if "%" in hdr else hdr[" %"]
    except KeyError as e:
        raise ValueError(f"{sheet!r} sheet is missing an expected header column: {e}. "
                         f"Found headers: {sorted(hdr)}") from None
    # on the combined sheet a `Batch` column sits between Name and %, so the
    # sessions still start after the percentage, wherever that has moved to
    first_session_col = pct_col + 1
    batch_col = hdr.get("batch")
    cols = _session_columns(ws, first_session_col)

    if excused not in ("exclude", "present"):
        raise ValueError(f"excused must be 'exclude' or 'present', not {excused!r}")

    per_session = {c: {"Y": 0, "N": 0, "E": 0} for c, _, _ in cols}
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        roll = ws.cell(row=r, column=roll_col).value
        if not roll:
            continue
        name = str(ws.cell(row=r, column=name_col).value or "").strip()
        batch = (str(ws.cell(row=r, column=batch_col).value or "").strip()
                 if batch_col else "")
        present = absent = exc = na = 0
        for c, _, _ in cols:
            v = ws.cell(row=r, column=c).value
            v = str(v).strip().upper() if v is not None else ""
            if v == "Y":
                present += 1
                per_session[c]["Y"] += 1
            elif v == "N":
                absent += 1
                per_session[c]["N"] += 1
            elif v == "E":
                exc += 1
                per_session[c]["E"] += 1
            elif v in NOT_APPLICABLE:
                na += 1
        credited = present + (exc if excused == "present" else 0)
        held = credited + absent
        rows.append(dict(Roll=str(roll).strip(), Name=name, Batch=batch, Present=present,
                         Absent=absent, Excused=exc, NotApplicable=na, Held=held,
                         Percent=(credited / held * 100) if held else None))

    students = pd.DataFrame(rows)

    sess_rows = []
    for c, d, label in sorted(cols, key=lambda t: t[1]):
        y, n, e = per_session[c]["Y"], per_session[c]["N"], per_session[c]["E"]
        credited = y + (e if excused == "present" else 0)
        held = credited + n
        sess_rows.append(dict(Session=label, Date=d.date(), Present=y, Absent=n,
                              Excused=e,
                              Percent=(credited / held * 100) if held else None))
    sessions = pd.DataFrame(sess_rows)
    return students, sessions


def percent_map(path: str, sheet: str,
                excused: str = EXCUSED_MODE) -> tuple[dict[str, dict], str]:
    """{roll: {percent, present, held, absent, excused}} plus an "as of" label, for
    printing each student's standing on the next class's signature sheet.

    The label names the date of the last session actually held, not today:
    a sheet printed on Friday for Monday's class still reports Wednesday's
    figures, and saying so stops the obvious "this is out of date" reply.
    Students with no session held yet get `percent=None` - they are carried
    in the map so the sheet can print "-" rather than leaving a blank cell
    that reads like an omission.
    """
    students, sessions = load(path, sheet, excused=excused)
    # Only sessions with marks in them. A column can carry a *future* date and
    # no marks yet (the term's remaining classes are pre-dated in this
    # workbook), and taking the plain max date would label the sheet with a
    # class that hasn't happened. An all-excused session still counts as held.
    held = (sessions[(sessions.Present + sessions.Absent + sessions.Excused) > 0]
            if len(sessions) else sessions)
    asof = ""
    if len(held):
        last = held.Date.max()
        asof = last.strftime("%d %b %Y") if hasattr(last, "strftime") else str(last)
    out = {}
    for r in students.itertuples():
        pct = None if pd.isna(r.Percent) else float(r.Percent)
        out[str(r.Roll).strip()] = dict(percent=pct, present=int(r.Present),
                                        held=int(r.Held), absent=int(r.Absent),
                                        excused=int(r.Excused),
                                        not_applicable=int(r.NotApplicable))
    return out, asof


# ─────────────────────── per-student marks, trends, forecast ───────────────────────
def load_marks(path: str, sheet: str) -> pd.DataFrame:
    """Every student's mark for every dated class, one row per (student, class).

    Columns: Roll, Name, Batch, Session, Date, Mark. `Mark` is `Y`, `N`, `E`,
    `NA` (a deliberate not-applicable) or `""` (blank - not marked, usually a
    class not yet held). `Batch` is the sheet's own Batch column on the combined
    sheet and the sheet name on a batch sheet, so a frame read from either
    groups the same way.

    `load` answers "how is this student doing"; this answers "which classes did
    they miss", which is the question behind every attendance query.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    ws = wb[sheet]
    hdr = _header_cols(ws)
    try:
        roll_col, name_col = hdr["roll number"], hdr["name"]
        pct_col = hdr["%"] if "%" in hdr else hdr[" %"]
    except KeyError as e:
        raise ValueError(f"{sheet!r} sheet is missing an expected header column: {e}. "
                         f"Found headers: {sorted(hdr)}") from None
    batch_col = hdr.get("batch")
    cols = sorted(_session_columns(ws, pct_col + 1), key=lambda t: t[1])
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        roll = ws.cell(row=r, column=roll_col).value
        if not roll:
            continue
        name = str(ws.cell(row=r, column=name_col).value or "").strip()
        batch = (str(ws.cell(row=r, column=batch_col).value or "").strip()
                 if batch_col else "") or sheet
        for c, d, label in cols:
            v = ws.cell(row=r, column=c).value
            v = str(v).strip().upper() if v is not None else ""
            mark = v if v in ("Y", "N", "E") else ("NA" if v in NOT_APPLICABLE else
                                                   ("" if not v else v))
            rows.append(dict(Roll=str(roll).strip(), Name=name, Batch=batch,
                             Session=label, Date=d.date(), Mark=mark))
    return pd.DataFrame(rows, columns=["Roll", "Name", "Batch", "Session", "Date", "Mark"])


def batch_trend(marks: pd.DataFrame, excused: str = EXCUSED_MODE) -> pd.DataFrame:
    """Turnout per class held, per batch, from `load_marks`.

    Columns: Batch, Date, Session, Present, Absent, Excused, Percent (that
    class's turnout) and Cumulative (the batch's attendance over every class up
    to and including this one - the figure a student's percentage is measured
    against). A class counts as held for a batch once anyone in it has a
    Y/N/E mark, so pre-dated future columns stay off the chart.
    """
    cols = ["Batch", "Date", "Session", "Present", "Absent", "Excused", "Percent", "Cumulative"]
    if marks is None or marks.empty:
        return pd.DataFrame(columns=cols)
    m = marks[marks.Mark.isin(["Y", "N", "E"])]
    if m.empty:
        return pd.DataFrame(columns=cols)
    g = (m.assign(Y=(m.Mark == "Y").astype(int), N=(m.Mark == "N").astype(int),
                  E=(m.Mark == "E").astype(int))
          .groupby(["Batch", "Date", "Session"], sort=True)[["Y", "N", "E"]].sum()
          .reset_index().sort_values(["Batch", "Date"]))
    credited = g.Y + (g.E if excused == "present" else 0)
    counted = credited + g.N
    g["Percent"] = (credited / counted.where(counted > 0)) * 100
    cum_c = credited.groupby(g.Batch).cumsum()
    cum_n = counted.groupby(g.Batch).cumsum()
    g["Cumulative"] = (cum_c / cum_n.where(cum_n > 0)) * 100
    g = g.rename(columns={"Y": "Present", "N": "Absent", "E": "Excused"})
    return g[cols].reset_index(drop=True)


def at_risk(students: pd.DataFrame, threshold: float, horizon: int,
            remaining: int | None = None) -> pd.DataFrame:
    """Students on or above `threshold` who would fall below it by missing
    the next `horizon` classes in a row.

    Uses `mailer.misses_allowed`, the counterpart of the number a shortfall
    notice quotes, so "can miss 2 more" here and "attend the next 6" in a mail
    come from the same arithmetic. `remaining` (classes still to be held)
    caps the horizon: nobody is at risk from classes that will never happen.
    Pass None when the course length isn't known.

    Adds CanMiss (classes they can still miss and stay on the threshold) and
    IfMissed (their percentage after missing `horizon` in a row). Sorted most
    exposed first.
    """
    import mailer  # local import: the one place this module needs mailer

    cols = list(students.columns) + ["CanMiss", "IfMissed"]
    horizon = int(horizon)
    if remaining is not None:
        horizon = min(horizon, max(0, int(remaining)))
    if students is None or students.empty or horizon <= 0:
        return pd.DataFrame(columns=cols)
    df = students[students.Held > 0].copy()
    df = df[df.Percent.notna() & (df.Percent >= float(threshold) - 1e-9)]
    if df.empty:
        return pd.DataFrame(columns=cols)
    credited = (df.Held - df.Absent).astype(int)
    df["CanMiss"] = [mailer.misses_allowed(c, int(a), threshold)
                     for c, a in zip(credited, df.Absent)]
    df["IfMissed"] = credited / (df.Held + horizon) * 100
    df = df[df.CanMiss < horizon]
    return df.sort_values(["CanMiss", "Percent"]).reset_index(drop=True)
