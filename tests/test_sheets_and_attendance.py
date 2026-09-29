"""Printed output and the register reader."""
import json
from pathlib import Path

import pandas as pd
import pytest

import config as C
import attendance_report as attrep
import seating
import sheets as sheetmod


def _alloc(n=12):
    seats = seating.seat_map("LHC110")
    seats = seats[(seats.Block == "A") & (seats.Kind == "core")].head(n)
    return pd.DataFrame({"Roll": [f"B26XX{i:04d}" for i in range(1, n + 1)],
                         "Name": [f"Student {i}" for i in range(1, n + 1)],
                         "Block": "A", "Seat": list(seats.Seat),
                         "SeatRow": list(seats.Row), "SeatCol": list(seats.Col)})


def test_sheet_rows_match_the_allocation_in_seat_order(tmp_path):
    a = _alloc()
    tpl = sheetmod.build_all(a, str(tmp_path), course="DEMO101", room="LHC110", session="S1")
    info = json.loads(open(tpl).read())
    rows = info["blocks"]["A"]["rows"]
    assert [r["roll"] for r in rows] == list(a.sort_values(["SeatRow", "SeatCol"]).Roll)
    assert not any("attendance" in r for r in rows), "plain sheets must carry no attendance column"


def test_attendance_column_is_opt_in_and_labelled(tmp_path):
    a = _alloc()
    att = {r: {"percent": 50.0, "present": 1, "held": 2} for r in a.Roll}
    tpl = sheetmod.build_all(a, str(tmp_path), course="DEMO101", room="LHC110", session="S1",
                             attendance=att, att_asof="24 Aug 2026")
    rows = json.loads(open(tpl).read())["blocks"]["A"]["rows"]
    assert all(r["attendance"] == "50%  1/2" for r in rows)


def test_missing_student_prints_a_dash_not_a_wrong_number():
    assert sheetmod._att_text(None) == "-"
    assert sheetmod._att_text({"percent": None, "held": 0}) == "-"


def test_stale_block_pdfs_are_cleared(tmp_path):
    a = _alloc()
    sheetmod.build_all(a, str(tmp_path), course="DEMO101", room="LHC110", session="S1")
    (tmp_path / "signature_block_Z.pdf").write_bytes(b"%PDF-1.4 stale")
    sheetmod.build_all(a, str(tmp_path), course="DEMO101", room="LHC110", session="S1")
    assert not (tmp_path / "signature_block_Z.pdf").exists(), \
        "a sheet for a block nobody sits in was left where it could be printed"


def test_percent_map_ignores_unmarked_future_classes():
    """The workbook pre-dates the whole term; 'as of' must be the last class
    that actually has marks, not the last dated column."""
    import pathlib
    wb = seating.ATTENDANCE_WORKBOOK
    if not pathlib.Path(wb).exists():
        import pytest
        pytest.skip("master workbook not present")
    m, asof = attrep.percent_map(wb, "English")
    students, sessions = attrep.load(wb, "English")
    marked = sessions[(sessions.Present + sessions.Absent) > 0]
    assert len(marked) < len(sessions), "expected some scheduled-but-unmarked columns"
    assert asof == marked.Date.max().strftime("%d %b %Y")
    assert set(m) == set(students.Roll)


def _excused_workbook(tmp_path):
    """A miniature Class Attendance workbook: header on row 3, dates on row 2,
    one Y/N/E column each."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "English"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "SR", "Roll Number", "Name", "%"
    from datetime import datetime
    for i, (col, label) in enumerate([("E", "C1"), ("F", "C2"), ("G", "C3")]):
        ws[f"{col}2"] = datetime(2026, 8, 5 + i)
        ws[f"{col}3"] = label
    rows = [("B26XX0001", "All present", "Y", "Y", "Y"),
            ("B26XX0002", "One excused", "Y", "E", "Y"),
            ("B26XX0003", "One absent", "Y", "N", "Y")]
    for r, (roll, name, *marks) in enumerate(rows, start=4):
        ws[f"B{r}"], ws[f"C{r}"] = roll, name
        for col, m in zip("EFG", marks):
            ws[f"{col}{r}"] = m
    path = tmp_path / "att.xlsx"
    wb.save(path)
    return str(path)


def test_excused_is_never_an_absence(tmp_path):
    students, _ = attrep.load(_excused_workbook(tmp_path), "English")
    by_roll = students.set_index("Roll")
    assert by_roll.at["B26XX0002", "Absent"] == 0, "an excused absence was counted as absent"
    assert by_roll.at["B26XX0002", "Excused"] == 1
    assert by_roll.at["B26XX0003", "Absent"] == 1


def test_excused_modes_change_only_the_percentage(tmp_path):
    wb = _excused_workbook(tmp_path)
    excluded = attrep.load(wb, "English")[0].set_index("Roll")
    credited = attrep.load(wb, "English", excused="present")[0].set_index("Roll")

    # dropped from the denominator: 2 of 2 count
    assert excluded.at["B26XX0002", "Held"] == 2
    assert excluded.at["B26XX0002", "Percent"] == 100.0
    # counted as attendance: 3 of 3
    assert credited.at["B26XX0002", "Held"] == 3
    assert credited.at["B26XX0002", "Percent"] == 100.0
    # the unexcused absence is unaffected by the mode
    assert (excluded.at["B26XX0003", "Percent"]
            == pytest.approx(credited.at["B26XX0003", "Percent"])
            == pytest.approx(200 / 3))


def test_excused_session_still_counts_as_held(tmp_path):
    """A class where everyone was excused still happened - it must not vanish
    from the 'as of' date."""
    m, asof = attrep.percent_map(_excused_workbook(tmp_path), "English")
    assert asof == "07 Aug 2026"
    assert m["B26XX0002"]["excused"] == 1 and m["B26XX0002"]["absent"] == 0


def test_shortfall_uses_unexcused_absences_only(tmp_path):
    import mailer
    students, _ = attrep.load(_excused_workbook(tmp_path), "English")
    students["Absent"] = [0, 0, 6]          # only the third has real absences
    pending = mailer.unmailed_attendance(students, path=tmp_path / "log.csv")
    assert list(pending.Roll) == ["B26XX0003"]


def test_not_applicable_is_distinct_from_blank(tmp_path):
    """A late joiner's pre-admission classes are marked N/A: counted towards
    nothing, but recorded as deliberate so they don't read as missing data."""
    import openpyxl
    from datetime import datetime
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "English"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "SR", "Roll Number", "Name", "%"
    for i, col in enumerate("EFG"):
        ws[f"{col}2"] = datetime(2026, 8, 5 + i)
        ws[f"{col}3"] = f"C{i + 1}"
    ws["B4"], ws["C4"] = "B22XX0001", "Late joiner"
    ws["E4"], ws["F4"], ws["G4"] = "N/A", "N/A", "Y"
    ws["B5"], ws["C5"] = "B26XX0002", "Blank cell"
    ws["G5"] = "Y"
    path = tmp_path / "att.xlsx"
    wb.save(path)

    students = attrep.load(str(path), "English")[0].set_index("Roll")
    late = students.loc["B22XX0001"]
    assert late.NotApplicable == 2 and late.Absent == 0 and late.Held == 1
    assert late.Percent == 100.0, "N/A must not drag a percentage down"
    assert students.loc["B26XX0002"].NotApplicable == 0, "a blank is not an N/A"


# ─────────────── the workbook's combined (both batches) sheet ───────────────
WB = C.ATTENDANCE_WORKBOOK


def _wb_present():
    return Path(WB).exists()


def test_combined_sheet_is_not_counted_as_a_third_cohort():
    """The workbook gained a sheet holding both batches on one list. It repeats
    the students on the two batch sheets, so anything that adds sheets together
    has to leave it out or count everyone twice."""
    import seating

    if not _wb_present():
        pytest.skip("register not present")

    assert attrep.is_combined(WB, "Combined")
    assert "Combined" not in attrep.cohort_sheets(WB)
    assert attrep.combined_sheets(WB) == ["Combined"]
    for sheet in attrep.cohort_sheets(WB):
        assert not attrep.is_combined(WB, sheet)

    # reading every sheet used to raise on duplicate roll numbers
    everyone = seating.load_attendance_roster(WB)
    assert not everyone.Roll.duplicated().any()
    assert set(everyone.Medium) == set(attrep.cohort_sheets(WB))

    # asked for by name it reads, taking each student's batch from the column
    combined = seating.load_attendance_roster(WB, "Combined")
    assert len(combined) == len(everyone)
    assert set(combined.Roll) == set(everyone.Roll)
    by_roll = dict(zip(everyone.Roll, everyone.Medium))
    for r in combined.itertuples():
        assert r.Medium == by_roll[r.Roll], f"{r.Roll} has the wrong batch on Combined"


def test_combined_sheet_marks_match_the_batch_sheets():
    """A combined sheet quietly out of step with the registers is worse than
    not having one, because it is the sheet people read."""
    import checks

    if not _wb_present():
        pytest.skip("register not present")

    per_sheet = {}
    for name in attrep.cohort_sheets(WB):
        stu, _ = attrep.load(WB, name)
        for r in stu.itertuples():
            per_sheet[str(r.Roll)] = (int(r.Present), int(r.Absent), int(r.Excused))

    comb, _ = attrep.load(WB, "Combined")
    assert len(comb) == len(per_sheet)
    for r in comb.itertuples():
        assert (int(r.Present), int(r.Absent), int(r.Excused)) == per_sheet[str(r.Roll)], \
            f"{r.Roll} differs between Combined and its batch sheet"

    results = checks.check_combined_sheet()
    assert results and all(r["status"] == "ok" for r in results), results


def test_summary_from_the_combined_sheet_matches_the_batch_sheets():
    """A summary read from the combined sheet has to give the same figures as
    reading the two batch sheets and adding them up - and it still says which
    batch each student is in, from the sheet's own Batch column."""
    if not _wb_present():
        pytest.skip("register not present")

    comb, comb_sessions = attrep.load(WB, "Combined")
    assert set(comb.Batch) == set(attrep.cohort_sheets(WB)), \
        "the combined sheet's Batch column does not name the batch sheets"

    parts = []
    for name in attrep.cohort_sheets(WB):
        stu, _ = attrep.load(WB, name)
        assert (stu.Batch == "").all(), "a batch sheet grew a Batch column"
        stu = stu.assign(Batch=name)
        parts.append(stu)
    by_sheets = pd.concat(parts, ignore_index=True).set_index("Roll").sort_index()
    by_combined = comb.set_index("Roll").sort_index()

    assert len(by_combined) == len(by_sheets)
    for col in ("Batch", "Present", "Absent", "Excused", "Held"):
        assert by_combined[col].equals(by_sheets[col]), f"{col} differs"

    # and the defaulter list is the same set of people either way
    def below(df, pct=75):
        seen = df[df.Held > 0]
        return set(seen[seen.Percent < pct].index)

    assert below(by_combined) == below(by_sheets)


def test_dashboard_sheet_is_read_not_summed_over():
    """The workbook carries the course's own live `Dashboard` tab. It holds no
    marks, so it must never be offered as a batch to read or mail from - but
    its settings are the course's, so they are read rather than a second copy
    of them being kept in the toolkit."""
    import checks

    if not _wb_present():
        pytest.skip("register not present")

    assert "Dashboard" in attrep.sheet_names(WB)
    assert not attrep.is_register(WB, "Dashboard")
    assert "Dashboard" not in attrep.register_sheets(WB)
    assert "Dashboard" not in attrep.cohort_sheets(WB)
    assert "Dashboard" not in attrep.combined_sheets(WB)
    assert attrep.register_sheets(WB) == attrep.cohort_sheets(WB) + attrep.combined_sheets(WB)

    blocks = attrep.dashboard_blocks(WB)
    assert float(blocks["settings"]["threshold (%)"]) == 75
    assert int(float(blocks["settings"]["total classes (course)"])) > 0
    assert set(blocks["columns"]) >= {"All"}
    assert blocks["bands"] and sum(blocks["bands"].values()) > 0

    # a sheet that is not there is not an error - this is a convenience
    assert attrep.dashboard_blocks(WB, sheet="No Such Sheet") == {}

    results = checks.check_dashboard_sheet()
    assert results and all(r["status"] in ("ok", "warn") for r in results)


def test_dashboard_sheet_figures_match_the_register():
    """What the course reads on that sheet has to be what the register says,
    including the count that decides which attendance mail a student gets."""
    import mailer

    if not _wb_present():
        pytest.skip("register not present")

    blocks = attrep.dashboard_blocks(WB)
    summary = blocks["summary"]
    threshold = float(blocks["settings"]["threshold (%)"])
    total = int(float(blocks["settings"]["total classes (course)"]))

    students, sessions = attrep.load(WB, "Combined")
    marked = sessions[(sessions.Present + sessions.Absent + sessions.Excused) > 0]
    remaining = total - len(marked)
    seen = students[students.Held > 0]
    below = seen[seen.Percent < threshold]

    assert int(summary["students"]["All"]) == len(students)
    assert int(summary["below threshold"]["All"]) == len(below)
    assert int(summary["critical (below 50%)"]["All"]) == int((seen.Percent < 50).sum())
    cannot = sum(1 for r in below.itertuples()
                 if not mailer.can_still_reach(int(r.Present), int(r.Absent),
                                               remaining, threshold))
    assert int(summary["can\'t reach threshold by end"]["All"]) == cannot

    # per batch too, where the sheet breaks it down
    for batch in ("English", "Hindi"):
        if batch not in blocks["columns"]:
            continue
        part = seen[seen.Batch == batch]
        assert int(summary["students"][batch]) == int((students.Batch == batch).sum())
        assert int(summary["below threshold"][batch]) == int((part.Percent < threshold).sum())


def test_attendance_report_pdf_states_the_figures_and_its_method(tmp_path):
    """The report is a drawing job over `attendance_report` and `mailer`, not a
    second calculation: a report whose numbers are worked out separately will
    eventually disagree with the mail a student receives."""
    import attendance_pdf as attpdf
    import mailer
    from pypdf import PdfReader

    if not _wb_present():
        pytest.skip("register not present")

    d = attpdf.report_data(WB)
    students, sessions = attrep.load(WB, d["sheet"])
    marked = sessions[(sessions.Present + sessions.Absent + sessions.Excused) > 0]
    assert d["held"] == len(marked)
    assert d["remaining"] == d["total"] - d["held"]
    # the split the two mail templates use has to be the same split here
    assert d["cannot_reach"] == sum(
        1 for r in d["below"].itertuples()
        if not mailer.can_still_reach(int(r.Present), int(r.Absent), d["remaining"],
                                      d["threshold"]))
    # a dated column nobody has marked is not a class held, and is called out
    assert all(g not in set(marked.Session) for g in d["gaps"])

    out = tmp_path / "report.pdf"
    attpdf.build_report(WB, str(out))
    pages = PdfReader(str(out)).pages
    assert 1 <= len(pages) <= 6
    text = " ".join((p.extract_text() or "") for p in pages)

    assert f"{len(d['students'])}" in text and "Students enrolled" in text
    assert "Batch comparison" in text and "Key observations" in text
    assert "Class-wise attendance" in text and "Method and notes" in text
    for r in d["below"].itertuples():           # everyone below the line is named
        assert r.Roll in text, f"{r.Roll} missing from the report"
    assert "Y / (Y + N)" in text, "the report does not state how the percentage is made"

    # switched off, those sections are gone rather than empty
    lean = tmp_path / "lean.pdf"
    attpdf.build_report(WB, str(lean), class_wise=False, method=False)
    lean_text = " ".join((p.extract_text() or "") for p in PdfReader(str(lean)).pages)
    assert "Class-wise attendance" not in lean_text and "Method and notes" not in lean_text
